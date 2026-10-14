"""Builds the DOCX (and diagram PNG) from the HTML/PDF source of truth.

Usage: python3 build_docx.py   (requires python-docx, lxml, pymupdf and the exported PDF)
"""
import re
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import html

HERE = Path(__file__).parent
NAME = "Zifamba-Tayanana_3220254_Cloud Programming_P1_S"
SRC = HERE / f"{NAME}.html"
PDF = HERE / f"{NAME}.pdf"
PNG = HERE / "architecture-diagram.png"
OUT = HERE / f"{NAME}.docx"

NAVY = RGBColor(0x23, 0x2F, 0x3E)
ORANGE = RGBColor(0xED, 0x71, 0x00)
GREY = RGBColor(0x6B, 0x72, 0x80)
LEAD_COLORS = {"p": RGBColor(0x6B, 0x2F, 0xD6), "b": RGBColor(0x14, 0x7E, 0xBA),
               "k": ORANGE, "g": RGBColor(0x3F, 0x7A, 0x0E)}
BODY_PT = 8.5


def export_diagram_png():
    doc = pymupdf.open(PDF)
    page = doc[0]
    top = page.search_for("3. AWS CLOUD ARCHITECTURE DIAGRAM")[0].y1 + 2
    bottom = page.search_for("Figure 1:")[0].y0 - 1
    clip = pymupdf.Rect(page.rect.x0 + 36, top, page.rect.x1 - 36, bottom)
    page.get_pixmap(dpi=300, clip=clip).save(PNG)


def set_font(run, size=BODY_PT, bold=None, italic=None, color=None):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def tight(par, after=1.5, before=0, line=1.15):
    pf = par.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing = line


def add_inline(par, el, size=BODY_PT, bold=False, italic=False, color=None):
    """Recursively copy text from an lxml element, keeping <b>/<i> styling."""
    if el.text:
        set_font(par.add_run(el.text), size, bold, italic, color)
    for child in el:
        cb, ci, cc = bold, italic, color
        if child.tag == "b":
            cb = True
            cc = LEAD_COLORS.get(child.get("class"), color)
        elif child.tag == "i":
            ci = True
        add_inline(par, child, size, cb, ci, cc)
        if child.tail:
            set_font(par.add_run(child.tail), size, bold, italic, color)


def normalise(el):
    for node in el.iter():
        if node.text:
            node.text = re.sub(r"\s+", " ", node.text)
        if node.tail is not None and node is not el:
            node.tail = re.sub(r"\s+", " ", node.tail)
    if el.text:
        el.text = el.text.lstrip()
    return el


def bottom_border(par, color="232F3E", size=12):
    p_pr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": str(size), "w:space": "1", "w:color": color}.items():
        b.set(qn(k), v)
    bdr.append(b)
    p_pr.append(bdr)


def heading(doc, text):
    par = doc.add_paragraph()
    tight(par, after=2, before=4)
    set_font(par.add_run("▌ "), 10.5, True, color=ORANGE)
    set_font(par.add_run(text.upper()), 10.5, True, color=NAVY)


def field(run, instr):
    for tag, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


def build():
    tree = html.parse(str(SRC)).getroot()
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin = Mm(11), Mm(10)
    sec.left_margin = sec.right_margin = Mm(15)
    sec.footer_distance = Mm(6)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(BODY_PT)
    doc.styles["Footer"].font.size = Pt(8)
    doc.styles["Footer"].font.color.rgb = GREY
    doc.styles["Footer"].paragraph_format.tab_stops.clear_all()

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col, cell, width in zip(table.columns, table.rows[0].cells, (Mm(52), Mm(46), Mm(82))):
        col.width = width
        cell.width = width
    for cell, (label, value), align in zip(
        table.rows[0].cells,
        [("Student Name: ", "Tayanana Zifamba"), ("Matriculation No: ", "3220254"),
         ("Course: ", "Cloud Programming (DLBSEPCP01_E)")],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT],
    ):
        par = cell.paragraphs[0]
        par.alignment = align
        tight(par, after=0)
        set_font(par.add_run(label), 8.5, True)
        set_font(par.add_run(value), 8.5, False, color=RGBColor(0x37, 0x41, 0x51))
    rule = doc.add_paragraph()
    tight(rule, after=4, line=0.6)
    bottom_border(rule)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(title, after=1)
    set_font(title.add_run("PORTFOLIO – PART 1: "), 13, True, color=NAVY)
    set_font(title.add_run("CONCEPTION PHASE"), 13, True, color=ORANGE)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(sub, after=2)
    set_font(sub.add_run("Task 1: Host a Highly Available, Globally Accessible Webpage on AWS"), 10, False, True,
             RGBColor(0x37, 0x41, 0x51))

    sections = tree.findall(".//section")

    heading(doc, "1. Concept & Objective")
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    tight(par, after=2)
    add_inline(par, normalise(sections[0].find(".//p")))

    heading(doc, "2. Architectural Decisions & Service Justifications")
    for li in sections[1].findall(".//li"):
        par = doc.add_paragraph(style="List Bullet")
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        tight(par, after=1.5)
        add_inline(par, normalise(li))

    heading(doc, "3. AWS Cloud Architecture Diagram")
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(pic, after=1, line=1.0)
    pic.add_run().add_picture(str(PNG), width=Mm(178))

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(cap, after=2)
    add_inline(cap, normalise(sections[2].find(".//figcaption")), size=7.5, color=RGBColor(0x4B, 0x55, 0x63))

    refs = doc.add_paragraph()
    tight(refs, after=0)
    add_inline(refs, normalise(sections[2].find(".//p")), size=7, color=GREY)

    foot = sec.footer.paragraphs[0]
    foot.text = ""
    tab_stops = foot.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Mm(180), alignment=2)
    set_font(foot.add_run("Submission Filename: Zifamba-Tayanana_3220254_Cloud Programming_P1_S.pdf\tpage "), 8,
             color=GREY)
    r = foot.add_run()
    set_font(r, 8, color=GREY)
    field(r, "PAGE")
    set_font(foot.add_run(" of "), 8, color=GREY)
    r = foot.add_run()
    set_font(r, 8, color=GREY)
    field(r, "NUMPAGES")

    doc.save(OUT)


if __name__ == "__main__":
    export_diagram_png()
    build()
    print(f"wrote {OUT.name} and {PNG.name}")
