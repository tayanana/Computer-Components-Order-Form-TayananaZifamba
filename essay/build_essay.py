"""Build the research essay as DOCX and PDF following the IU formatting rules.

Usage: python3 build_essay.py
Requires python-docx, LibreOffice (soffice) and poppler-utils (pdftotext).
"""

import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).parent
OUT_NAME = "Research_Essay_Zifamba_ISSE01_E"
FONT = "Arial"
TEXT_WIDTH_CM = 21.0 - 2 * 2.0

TITLE_PAGE = {
    "title": "Institutionalizing Clean Code: Strategies for Organizational Implementation",
    "fields": [
        ("Type of Thesis", "Research Essay"),
        ("Course Name", "ISSE01_E Seminar: Software Engineering"),
        ("Course of Study", "Bachelor Software Development"),
        ("Date", "October 2026"),
        ("Author", "Tayanana Zifamba"),
        ("Matriculation Number", "[Insert Number]"),
        ("Tutor's Name", "Prof. Dr. Carsten Skerra"),
    ],
}

ABBREVIATIONS = [
    ("AI", "Artificial Intelligence"),
    ("API", "Application Programming Interface"),
    ("CI", "Continuous Integration"),
    ("CI/CD", "Continuous Integration and Continuous Delivery"),
]

ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"]


def parse_markdown(path):
    """Return a list of (kind, text) blocks, kind in h1/h2/h3/p."""
    blocks = []
    for chunk in path.read_text(encoding="utf-8").split("\n\n"):
        chunk = chunk.strip()
        if not chunk:
            continue
        m = re.match(r"^(#{1,3}) (.+)$", chunk)
        if m:
            blocks.append((f"h{len(m.group(1))}", m.group(2).strip()))
        else:
            blocks.append(("p", " ".join(chunk.split())))
    return blocks


def set_run_font(run, size, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)


def add_rich_text(par, text, size=11, bold=False):
    """Add text where *segments* are italic."""
    for i, part in enumerate(re.split(r"\*(.+?)\*", text)):
        if part:
            set_run_font(par.add_run(part), size, bold=bold, italic=(i % 2 == 1))


def format_paragraph(par, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, spacing=1.5):
    pf = par.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = spacing


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    format_paragraph_style(normal.paragraph_format, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 6)
    for level in (1, 2, 3):
        st = doc.styles[f"Heading {level}"]
        st.font.name = FONT
        st.font.size = Pt(12)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = None
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), FONT)
        for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rfonts.get(qn(theme_attr)) is not None:
                del rfonts.attrib[qn(theme_attr)]
        format_paragraph_style(st.paragraph_format, WD_ALIGN_PARAGRAPH.LEFT, 12, 6)
        st.paragraph_format.keep_with_next = True


def format_paragraph_style(pf, align, before, after):
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.5


def set_margins(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, side, Cm(2.0))
    section.footer_distance = Cm(1.0)


def set_page_numbering(section, fmt, start=1):
    sect_pr = section._sectPr
    pg = sect_pr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sect_pr.append(pg)
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), str(start))


def add_page_field(par):
    run = par.add_run()
    set_run_font(run, 11)
    for tag, text in (("begin", None), (None, "PAGE"), ("separate", None), (None, "1"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        elif text == "PAGE":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = " PAGE "
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._element.append(el)


def centered_footer(footer):
    footer.is_linked_to_previous = False
    par = footer.paragraphs[0]
    par.text = ""
    format_paragraph(par, WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 1.0)
    add_page_field(par)


def add_heading(doc, text, level, page_break_before=False):
    h = doc.add_heading(level=level)
    add_rich_text(h, text, size=12, bold=True)
    h.paragraph_format.page_break_before = page_break_before
    return h


def build_title_page(doc):
    def line(text, size=11, bold=False, after=6, align=WD_ALIGN_PARAGRAPH.CENTER):
        p = doc.add_paragraph()
        format_paragraph(p, align, 0, after)
        add_rich_text(p, text, size=size, bold=bold)
        return p

    line("IU International University of Applied Sciences", 12, True, 72)
    line(TITLE_PAGE["title"], 14, True, 48)
    for label, value in TITLE_PAGE["fields"]:
        p = doc.add_paragraph()
        format_paragraph(p, WD_ALIGN_PARAGRAPH.LEFT, 0, 6)
        p.paragraph_format.left_indent = Cm(3.0)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(8.0))
        set_run_font(p.add_run(f"{label}:"), 11, bold=True)
        set_run_font(p.add_run(f"\t{value}"), 11)


def build_toc(doc, entries, pages):
    title = doc.add_paragraph()
    format_paragraph(title, WD_ALIGN_PARAGRAPH.LEFT, 0, 12)
    add_rich_text(title, "Table of Contents", size=12, bold=True)
    title.paragraph_format.page_break_before = True
    for level, text in entries:
        p = doc.add_paragraph()
        format_paragraph(p, WD_ALIGN_PARAGRAPH.LEFT, 0, 6)
        p.paragraph_format.left_indent = Cm(0.6 * (level - 1))
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        add_rich_text(p, f"{text}\t{pages.get(text, '0')}", size=11, bold=(level == 1))


def build_abbreviations(doc):
    title = doc.add_paragraph()
    format_paragraph(title, WD_ALIGN_PARAGRAPH.LEFT, 0, 12)
    add_rich_text(title, "List of Abbreviations", size=12, bold=True)
    title.paragraph_format.page_break_before = True
    for abbr, meaning in ABBREVIATIONS:
        p = doc.add_paragraph()
        format_paragraph(p, WD_ALIGN_PARAGRAPH.LEFT, 0, 6)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(3.0))
        add_rich_text(p, f"{abbr}\t{meaning}")


def build_document(blocks, references, pages, out_path):
    doc = Document()
    setup_styles(doc)

    front = doc.sections[0]
    set_margins(front)
    front.different_first_page_header_footer = True
    set_page_numbering(front, "upperRoman", 1)
    centered_footer(front.footer)

    build_title_page(doc)
    toc_entries = [(1, "List of Abbreviations")]
    toc_entries += [(int(k[1]), t) for k, t in blocks if k.startswith("h")]
    toc_entries.append((1, "References"))
    build_toc(doc, toc_entries, pages)
    build_abbreviations(doc)

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    set_margins(body)
    body.different_first_page_header_footer = False
    set_page_numbering(body, "decimal", 1)
    centered_footer(body.footer)

    for kind, text in blocks:
        if kind == "p":
            p = doc.add_paragraph()
            add_rich_text(p, text)
        else:
            add_heading(doc, text, int(kind[1]))

    add_heading(doc, "References", 1, page_break_before=True)
    for ref in references:
        p = doc.add_paragraph()
        format_paragraph(p, WD_ALIGN_PARAGRAPH.LEFT, 0, 6)
        p.paragraph_format.left_indent = Cm(1.27)
        p.paragraph_format.first_line_indent = Cm(-1.27)
        add_rich_text(p, ref)

    doc.save(out_path)


def to_pdf(docx_path):
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
        check=True,
        capture_output=True,
    )
    return docx_path.with_suffix(".pdf")


def pdf_pages(pdf_path):
    text = subprocess.run(["pdftotext", "-layout", str(pdf_path), "-"], check=True, capture_output=True, text=True).stdout
    return [" ".join(p.split()) for p in text.split("\f")]


def locate_pages(pdf_path, blocks):
    pages = pdf_pages(pdf_path)
    first_sentence = blocks[next(i for i, (k, _) in enumerate(blocks) if k == "p")][1][:60]
    body_start = next(i for i, p in enumerate(pages) if first_sentence in p)
    result = {"List of Abbreviations": ROMAN[next(i for i, p in enumerate(pages) if p.startswith("List of Abbreviations"))]}
    search_from = body_start
    for kind, text in [b for b in blocks if b[0].startswith("h")] + [("h1", "References")]:
        plain = text.replace("*", "")
        idx = next(i for i in range(search_from, len(pages)) if plain in pages[i])
        result[text] = str(idx - body_start + 1)
        search_from = idx
    return result, pages, body_start


def main():
    blocks = parse_markdown(HERE / "essay.md")
    references = [" ".join(r.split()) for r in (HERE / "references.md").read_text(encoding="utf-8").split("\n\n") if r.strip()]
    docx_path = HERE / f"{OUT_NAME}.docx"

    pages = {}
    for _ in range(2):
        build_document(blocks, references, pages, docx_path)
        pdf_path = to_pdf(docx_path)
        pages, pdf_text, body_start = locate_pages(pdf_path, blocks)
    build_document(blocks, references, pages, docx_path)
    pdf_path = to_pdf(docx_path)
    pages, pdf_text, body_start = locate_pages(pdf_path, blocks)

    ref_page = int(pages["References"])
    core_pages = ref_page - 1
    words = {"intro": 0, "body": 0, "conclusion": 0}
    part = None
    for kind, text in blocks:
        if kind == "h1":
            part = {"1": "intro", "2": "body", "3": "conclusion"}[text[0]]
        elif kind == "p":
            words[part] += len(text.split())
    total = sum(words.values())
    print(f"PDF pages: {len(pdf_text)}, front matter pages: {body_start}, core text pages: {core_pages}")
    for k, v in words.items():
        print(f"  {k:<11} {v:>5} words  {100 * v / total:5.1f}%")
    print(f"  total       {total:>5} words")
    print("TOC:", pages)


if __name__ == "__main__":
    main()
